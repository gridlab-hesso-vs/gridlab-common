"""
Word order, offset and single-register writes.

The defaults keep every existing register map decoding and encoding exactly as
before; the options cover layouts such as the CNTE EMS protocol (32-bit values
with the low word first, actual = raw x scale + offset, FC06 writes only).

A never-started ModbusTCPServer is the interface: its read/write methods operate
directly on the in-memory datastore, so no socket is involved.
"""

import openpyxl
import pytest

from gridlab_modbus.modbus_data_model import Block, ModbusData, VarSpec
from gridlab_modbus.modbus_tcp import ModbusTCPClient, ModbusTCPServer
from gridlab_modbus.modbus_xls_parser import parse_modbusdata_from_xls

HOLDING = 3


def store() -> ModbusTCPServer:
    return ModbusTCPServer(host="127.0.0.1", port=0)


def words(server: ModbusTCPServer, address: int, count: int) -> list[int]:
    return list(server._context[0].getValues(HOLDING, address, count))


def block_of(server: ModbusTCPServer, direction: str, *specs: VarSpec, address: int = 100) -> Block:
    block = Block(storage="holding_register", initial_address=address, direction=direction)
    for spec in specs:
        block.add_var(spec)
    block.assign_modbus_interface(server)
    return block


# ---- word order ----

def test_default_word_order_is_high_word_first():
    srv = store()
    srv.write_holding_registers(100, [0x0001, 0x1170])            # 70000 = 0x00011170
    block = block_of(srv, "input", VarSpec(name="e", value_type="uint32"))
    block.sync_read()
    assert block.var_values["e"] == 70000


def test_little_word_order_decodes_low_word_first():
    srv = store()
    neg = -12345 & 0xFFFFFFFF
    srv.write_holding_registers(100, [0x1170, 0x0001, neg & 0xFFFF, neg >> 16, 0x0000, 0x3FC0])
    block = block_of(
        srv, "input",
        VarSpec(name="u", value_type="uint32", word_order="little"),
        VarSpec(name="s", value_type="sint32", word_order="little", scale_factor=0.1, decimal_digits=1),
        VarSpec(name="f", value_type="float", word_order="little", decimal_digits=1),   # 1.5 = 0x3FC00000
    )
    block.sync_read()
    assert block.var_values == {"u": 70000, "s": -1234.5, "f": 1.5}


def test_little_word_order_encodes_low_word_first():
    srv = store()
    block = block_of(srv, "output", VarSpec(name="p", value_type="sint32", word_order="little", scale_factor=0.1), address=300)
    block.set_value("p", -1234.5)
    block.sync_write()
    neg = -12345 & 0xFFFFFFFF
    assert words(srv, 300, 2) == [neg & 0xFFFF, neg >> 16]


def test_word_order_must_be_big_or_little():
    with pytest.raises(ValueError, match="word_order"):
        VarSpec(name="x", value_type="uint32", word_order="middle")


# ---- offset ----

def test_offset_applies_after_scale_and_is_inverted_on_write():
    srv = store()
    srv.write_holding_registers(100, [20123])
    block = block_of(srv, "input", VarSpec(name="i", value_type="uint16", scale_factor=0.1, offset=-2000, decimal_digits=1))
    block.sync_read()
    assert block.var_values["i"] == 12.3

    out = block_of(srv, "output", VarSpec(name="t", value_type="sint16", offset=-50), address=200)
    out.set_value("t", 25)
    out.sync_write()
    assert words(srv, 200, 1) == [75]


def test_integer_encode_drops_float_noise_but_still_truncates():
    srv = store()
    block = block_of(
        srv, "output",
        VarSpec(name="a", value_type="uint16", scale_factor=0.1),
        VarSpec(name="b", value_type="uint16"),
        address=500,
    )
    block.set_value("a", 0.3)      # 0.3 / 0.1 = 2.9999999999999996
    block.set_value("b", 12.9)     # a genuinely fractional value still truncates, as before
    block.sync_write()
    assert words(srv, 500, 2) == [3, 12]


# ---- writing one variable ----

def test_write_var_value_leaves_the_rest_of_the_block_alone():
    srv = store()
    srv.write_holding_registers(400, [7, 8, 9])
    data = ModbusData()
    block = Block(storage="holding_register", initial_address=400, direction="output")
    block.add_var(VarSpec(name="power_command", value_type="uint16"))
    block.add_var(VarSpec(name="setpoint", value_type="sint16", scale_factor=0.1))
    block.add_var(VarSpec(name="fault_reset", value_type="uint16"))
    data.add_block(block)
    data.assign_modbus_interface(srv)

    data.write_var_value("setpoint", -80.0)
    assert words(srv, 400, 3) == [7, -800 & 0xFFFF, 9]
    assert data.get_var_value("setpoint") == -80.0
    assert data.stats.write_ok == 1
    with pytest.raises(KeyError):
        data.write_var_value("not_a_variable", 1)


def test_write_var_value_on_a_coil_block_writes_that_coil_only():
    srv = store()
    srv.write_coils(10, [True, False, True])
    data = ModbusData()
    block = Block(storage="coil", initial_address=10, direction="output")
    for name in ("a", "b", "c"):
        block.add_bool(name)
    data.add_block(block)
    data.assign_modbus_interface(srv)

    data.write_var_value("b", True)
    data.write_var_value("a", False)
    assert [bool(x) for x in srv._context[0].getValues(1, 10, 3)] == [False, True, True]


# ---- single-register writes (FC06) ----

class _RecordingPymodbusClient:
    """Stands in for pymodbus' ModbusTcpClient: records the function code of each write."""

    class _Ok:
        def isError(self):
            return False

    def __init__(self):
        self.calls = []

    def write_register(self, address, value, slave=1):
        self.calls.append(("fc06", address, value))
        return self._Ok()

    def write_registers(self, address, values, slave=1):
        self.calls.append(("fc16", address, list(values)))
        return self._Ok()


@pytest.mark.parametrize("options, expected", [
    ({}, [("fc16", 10, [1, 2])]),                                            # default: unchanged
    ({"single_register_writes": True}, [("fc06", 10, 1), ("fc06", 11, 2)]),
])
def test_client_write_function_code(options, expected):
    client = ModbusTCPClient(host="127.0.0.1", port=1, **options)
    client._client = _RecordingPymodbusClient()
    client.write_holding_registers(10, [1, 2], slave=1)
    assert client._client.calls == expected


def test_client_write_holding_register_is_fc06():
    client = ModbusTCPClient(host="127.0.0.1", port=1)
    client._client = _RecordingPymodbusClient()
    client.write_holding_register(10, 5, slave=1)
    assert client._client.calls == [("fc06", 10, 5)]


# ---- spreadsheet columns ----

HEADER = ["Address", "Name", "Direction", "Type", "Size (bytes)", "Scale factor", "Unit",
          "Default value", "Min value", "Max value", "Decimal digits", "Notes"]


def _xlsx(path, header, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "holding registers"
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.create_sheet("input registers").append(
        ["Address", "Name", "Type", "Size (bytes)", "Scale factor", "Unit", "Default value",
         "Min value", "Max value", "Decimal digits", "Notes"])
    wb.create_sheet("coils").append(["Address", "Name", "Direction", "Default value", "Notes"])
    wb.create_sheet("discrete inputs").append(["Address", "Name", "Default value", "Notes"])
    wb.save(path)
    return path


def test_offset_and_word_order_columns_are_optional(tmp_path):
    plain = _xlsx(tmp_path / "plain.xlsx", HEADER, [[0, "p", "R", "sint32", None, 0.1, "kW", 0, None, None, 1, None]])
    (spec,) = [v for _, v in parse_modbusdata_from_xls(plain).get_var_specs()]
    assert (spec.offset, spec.word_order) == (0, "big")

    new = _xlsx(tmp_path / "new.xlsx", HEADER + ["Offset", "Word order"],
                [[0, "p", "R", "sint32", None, 0.1, "kW", 0, None, None, 1, None, -2000, "low first"]])
    (spec,) = [v for _, v in parse_modbusdata_from_xls(new).get_var_specs()]
    assert (spec.offset, spec.word_order) == (-2000, "little")


def test_unknown_word_order_cell_is_rejected(tmp_path):
    bad = _xlsx(tmp_path / "bad.xlsx", HEADER + ["Word order"],
                [[0, "p", "R", "sint32", None, 1, "", 0, None, None, 0, None, "sideways"]])
    with pytest.raises(ValueError, match="Word order"):
        parse_modbusdata_from_xls(bad)
